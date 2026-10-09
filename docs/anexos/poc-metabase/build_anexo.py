"""Genera el DOCX del anexo de la POC de Metabase con las capturas incluidas."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
CAP = HERE / "capturas"
OUT = Path.home() / "Downloads" / "Cardio_Insights_Anexo_POC_Metabase.docx"

SQL = """SELECT CAST(YEAR(FechaRealizado) AS CHAR) AS anio,
       CASE CodSeguroCoo
            WHEN 10 THEN 'FNR'
            WHEN 20 THEN 'Particular'
            ELSE 'Otro'
       END AS financiador,
       COUNT(*) AS actos
FROM flow_coordina
GROUP BY anio, financiador
ORDER BY anio, financiador"""

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)


def par(text, italic=False, bold=False, size=11, after=8):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    r.bold = bold
    r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(after)
    return p


def code(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    pf = p.paragraph_format
    pf.left_indent = Cm(0.8)
    pf.space_before = Pt(6)
    pf.space_after = Pt(10)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F4F6F8")
    pPr.append(shd)
    return p


def figure(name, text):
    doc.add_picture(str(CAP / name), width=Cm(16.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)


def head(text, level):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.color.rgb = RGBColor(0x0D, 0x3B, 0x66)
    return p


head("Anexo 12.5 – Prueba de concepto de integración con Metabase", 1)

head("Objetivo", 2)
par(
    "La prueba de concepto tuvo como objetivo validar técnicamente que Metabase podía cumplir el "
    "rol de componente de visualización dentro de Cardio Insights antes de comprometer la decisión "
    "en la arquitectura. Se buscó verificar cuatro aspectos considerados críticos: la instalación "
    "y ejecución de la herramienta, la conexión con una base de datos MySQL, la construcción de "
    "consultas y tableros sobre esa base, y la posibilidad de embeber un tablero dentro de una "
    "aplicación web propia mediante un acceso controlado."
)

head("Entorno utilizado", 2)
par(
    "La prueba se realizó sobre un entorno local, independiente del entorno institucional. Metabase "
    "se ejecutó como contenedor a partir de la imagen oficial, en el puerto 3001, y se conectó a "
    "una instancia de MySQL 8.0 con una base de prueba del esquema incc provista para el "
    "desarrollo. La base utilizada contiene un conjunto reducido de registros de prueba, lo que "
    "resultó suficiente para validar la integración: el objetivo de esta instancia no era analizar "
    "resultados clínicos, sino verificar el comportamiento técnico de la herramienta."
)

head("Consulta utilizada", 2)
par(
    "Para la prueba se definió una consulta simple pero representativa de los indicadores del "
    "producto, equivalente al requerimiento RF02: la cantidad de actos realizados por año, "
    "diferenciando el destino de facturación entre FNR y particulares."
)
code(SQL)

head("Procedimiento y resultados", 2)
par(
    "La consulta se escribió en el editor SQL de Metabase, sobre la conexión creada hacia la base "
    "de prueba. La herramienta permitió ejecutar la sentencia y obtener el resultado tabular, "
    "reconociendo las tablas del esquema conectado."
)
figure(
    "poc-01-consulta-sql.png",
    "Figura A.1 – Ejecución de la consulta en el editor SQL de Metabase, con el resultado obtenido "
    "y las tablas del esquema conectado.",
)

par(
    "A partir del mismo resultado se validó la generación de visualizaciones. La consulta se guardó "
    "como pregunta y se representó como gráfico de barras apiladas, definiendo el año como "
    "dimensión, el financiador como serie y la cantidad de actos como métrica. Esto permitió "
    "confirmar que la herramienta construye la visualización a partir de la consulta, sin necesidad "
    "de implementar la lógica de graficado en la aplicación."
)
figure(
    "poc-02-visualizacion.png",
    "Figura A.2 – La misma consulta guardada como pregunta y representada como gráfico de barras "
    "apiladas.",
)

par(
    "Luego se creó un tablero de ejemplo que agrupa dos representaciones de la consulta, una "
    "gráfica y una tabular, con el fin de verificar la organización de varios elementos dentro de "
    "una misma pantalla, tal como se prevé para los tableros de indicadores del producto."
)
figure(
    "poc-03-dashboard-metabase.png",
    "Figura A.3 – Tablero de ejemplo construido en Metabase con dos representaciones de la misma "
    "consulta.",
)

par(
    "Finalmente se validó el aspecto más relevante para la arquitectura: el embebido del tablero "
    "dentro de una aplicación web propia. Para ello se habilitó el embebido estático en Metabase y "
    "se generó una dirección firmada con la clave de la instancia, indicando el tablero autorizado "
    "y un tiempo de vigencia. Esa dirección se utilizó desde una página de la aplicación, que "
    "muestra el tablero dentro de su propia navegación. La prueba confirmó que el usuario accede al "
    "tablero sin ingresar a Metabase y sin conocer sus credenciales, ya que la clave de firma "
    "permanece del lado del servidor."
)
figure(
    "poc-04-embebido-en-la-app.png",
    "Figura A.4 – Tablero embebido dentro de una pantalla de Cardio Insights mediante una dirección "
    "firmada de vigencia limitada.",
)

head("Conclusiones de la prueba", 2)
par(
    "La prueba de concepto permitió confirmar que Metabase resultaba viable como componente de "
    "visualización para Cardio Insights. En particular, se verificó que la herramienta se instala y "
    "ejecuta mediante contenedores sin configuración compleja, que se conecta a MySQL sin requerir "
    "modificaciones en la base, que permite construir consultas, visualizaciones y tableros sobre "
    "esa conexión, y que un tablero puede visualizarse dentro de una aplicación propia mediante un "
    "acceso firmado y de vigencia limitada."
)
par(
    "Sobre esa base, el equipo consideró viable avanzar con Metabase en la integración definitiva, "
    "incorporando la generación de la dirección firmada en el backend, la protección de las rutas "
    "de la aplicación y la conexión de la herramienta con un usuario de base de datos con permisos "
    "de solo lectura, tal como se describe en el capítulo de arquitectura."
)

head("Limitaciones de la prueba", 2)
par(
    "La prueba se realizó en un entorno local y con un conjunto reducido de datos de prueba, por lo "
    "que no permite extraer conclusiones sobre el rendimiento con el volumen histórico completo de "
    "la institución ni sobre el comportamiento de la herramienta en el entorno institucional "
    "definitivo. La validación de la firma de acceso se realizó de forma manual durante la prueba; "
    "su generación desde el backend, junto con la verificación del rol del usuario, se implementó "
    "posteriormente como parte del producto."
)

props = doc.core_properties
props.author = "Equipo Cardio Insights"
props.last_modified_by = "Equipo Cardio Insights"
props.title = "Anexo - POC de integración con Metabase"
props.comments = ""

doc.save(OUT)
print("OK ->", OUT)
